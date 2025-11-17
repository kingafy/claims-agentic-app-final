import os
import io
import asyncio
import logging
from typing import Dict, List

import streamlit as st

from google.adk.runners import InMemoryRunner
from google.genai import types
from langfuse import get_client

from agents.claims_adk_app.agent import root_agent, regeneration_agent, qa_agent
from dotenv import load_dotenv
from google import genai
load_dotenv()
api_key=os.environ.get("GEMINI_API_KEY")
if not api_key:
    raise ValueError("GEMINI_API_KEY environment variable not set.")
else:
    print(api_key)
QA_CLIENT =genai.Client(api_key=api_key)

logging.basicConfig(level=logging.ERROR)

APP_NAME = "claims_adk_app"
USER_ID = "streamlit_user"

PIPELINE_RUNNER_KEY = "pipeline_runner"
QA_RUNNER_KEY = "qa_runner"
REGEN_RUNNER_KEY = "regen_runner"

UPLOADED_DOC_KEY = "uploaded_doc_text"
PIPELINE_RESULT_KEY = "pipeline_result"
QA_RESULT_KEY = "qa_last_answer"

LANGFUSE_CLIENT_KEY = "langfuse_client"

STEP_LABELS: Dict[str, str] = {
    "DataIngestionAgent": "Data Ingestion Agent",
    "ParserAgent": "Parser Agent",
    "SummaryAgent": "Summary Agent",
    "FormatterAgent": "Formatter Agent",
    "CriticAgent": "Critic Agent",
    "RegenerationAgent": "Regeneration Agent",
}


def require_api_key() -> None:
    api_key = os.getenv("GOOGLE_API_KEY") or os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError(
            "No Gemini API key found. Please set GOOGLE_API_KEY or GEMINI_API_KEY."
        )


def get_pipeline_runner() -> InMemoryRunner:
    if PIPELINE_RUNNER_KEY not in st.session_state:
        st.session_state[PIPELINE_RUNNER_KEY] = InMemoryRunner(
            agent=root_agent, app_name=APP_NAME
        )
    return st.session_state[PIPELINE_RUNNER_KEY]


def get_qa_runner() -> InMemoryRunner:
    if QA_RUNNER_KEY not in st.session_state:
        st.session_state[QA_RUNNER_KEY] = InMemoryRunner(
            agent=qa_agent, app_name=APP_NAME
        )
    return st.session_state[QA_RUNNER_KEY]


def get_regen_runner() -> InMemoryRunner:
    if REGEN_RUNNER_KEY not in st.session_state:
        st.session_state[REGEN_RUNNER_KEY] = InMemoryRunner(
            agent=regeneration_agent, app_name=APP_NAME
        )
    return st.session_state[REGEN_RUNNER_KEY]


def get_langfuse():
    if LANGFUSE_CLIENT_KEY in st.session_state:
        return st.session_state[LANGFUSE_CLIENT_KEY]

    try:
        client = get_client()
    except Exception:
        st.session_state[LANGFUSE_CLIENT_KEY] = None
        return None

    st.session_state[LANGFUSE_CLIENT_KEY] = client
    return client


def parse_uploaded_file(uploaded_file) -> str:
    if uploaded_file is None:
        return ""

    file_bytes = uploaded_file.read()
    mime = uploaded_file.type or ""
    name = uploaded_file.name.lower()

    if "pdf" in mime or name.endswith(".pdf"):
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            texts = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                texts.append(page_text)
            return "\n".join(texts).strip()
        except Exception as e:
            return f"[Error reading PDF: {e}]"

    if "word" in mime or name.endswith(".docx"):
        try:
            import docx

            document = docx.Document(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in document.paragraphs).strip()
        except Exception as e:
            return f"[Error reading DOCX: {e}]"

    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        return f"[Error reading file as text: {e}]"


async def run_pipeline_async(
    claim_text: str,
    status_placeholders: Dict[str, st.delta_generator.DeltaGenerator],
) -> str:
    runner = get_pipeline_runner()
    langfuse = get_langfuse()

    async def _run_core() -> str:
        session = await runner.session_service.create_session(
            app_name=APP_NAME,
            user_id=USER_ID,
            state={"claim_text": claim_text},
        )

        content = types.Content(
            role="user",
            parts=[types.Part(text="Process the uploaded claim and produce a final summary.")],
        )

        final_summary = "Pipeline did not return a summary."
        completed_steps: List[str] = []

        async for event in runner.run_async(
            user_id=USER_ID,
            session_id=session.id,
            new_message=content,
        ):
            author = getattr(event, "author", "")

            if author in STEP_LABELS and author not in completed_steps:
                label = STEP_LABELS.get(author, author)
                status_placeholders[author].success(f"✅ {label}: complete")
                completed_steps.append(author)

            if getattr(event, "is_final_response", None) and event.is_final_response():
                if event.content and event.content.parts:
                    text_parts = [
                        p.text for p in event.content.parts if getattr(p, "text", None)
                    ]
                    if text_parts:
                        final_summary = "\n".join(text_parts)

        return final_summary

    if langfuse:
        with langfuse.start_as_current_observation(
            as_type="span",
            name="claims_pipeline",
            input={"claim_excerpt": claim_text[:1000]},
            metadata={"app": APP_NAME, "mode": "pipeline", "user_id": USER_ID},
        ) as span:
            final_summary = await _run_core()
            span.update(output=final_summary[:4000])
            return final_summary

    return await _run_core()


def run_pipeline_sync(
    claim_text: str,
    status_placeholders: Dict[str, st.delta_generator.DeltaGenerator],
) -> str:
    return asyncio.run(run_pipeline_async(claim_text, status_placeholders))


async def run_qa_async(question: str, claim_text: str) -> str:
    runner = get_qa_runner()
    langfuse = get_langfuse()

    async def _run_core() -> str:
        session = await runner.session_service.create_session(
            app_name=APP_NAME,
            user_id=USER_ID,
            state={"claim_text": claim_text, "user_question": question},
        )

        content = types.Content(
            role="user",
            parts=[types.Part(text=f"User question: {question}")],
        )

        answer = "I couldn't generate an answer."
        async for event in runner.run_async(
            user_id=USER_ID,
            session_id=session.id,
            new_message=content,
        ):
            if getattr(event, "is_final_response", None) and event.is_final_response():
                if event.content and event.content.parts:
                    text_parts = [
                        p.text for p in event.content.parts if getattr(p, "text", None)
                    ]
                    if text_parts:
                        answer = "\n".join(text_parts)

        return answer

    if langfuse:
        with langfuse.start_as_current_observation(
            as_type="span",
            name="claim_qa",
            input={"question": question, "claim_excerpt": claim_text[:1000]},
            metadata={"app": APP_NAME, "mode": "qa", "user_id": USER_ID},
        ) as span:
            answer = await _run_core()
            span.update(output=answer[:4000])
            return answer

    return await _run_core()


def run_qa_sync(question: str, claim_text: str) -> str:
    return asyncio.run(run_qa_async(question, claim_text))


async def run_regen_async(existing_summary: str, feedback: str) -> str:
    runner = get_regen_runner()
    langfuse = get_langfuse()

    critique = feedback.strip()
    if "VERDICT:" not in critique.upper():
        critique = critique + "\n\nVERDICT: revise"

    async def _run_core() -> str:
        session = await runner.session_service.create_session(
            app_name=APP_NAME,
            user_id=USER_ID,
            state={
                "formatted_summary": existing_summary,
                "critique": critique,
            },
        )

        content = types.Content(
            role="user",
            parts=[types.Part(text="Regenerate the claim summary based on the critique.")],
        )

        new_summary = existing_summary
        async for event in runner.run_async(
            user_id=USER_ID,
            session_id=session.id,
            new_message=content,
        ):
            if getattr(event, "is_final_response", None) and event.is_final_response():
                if event.content and event.content.parts:
                    text_parts = [
                        p.text for p in event.content.parts if getattr(p, "text", None)
                    ]
                    if text_parts:
                        new_summary = "\n".join(text_parts)

        return new_summary

    if langfuse:
        with langfuse.start_as_current_observation(
            as_type="span",
            name="summary_regeneration",
            input={
                "existing_summary": existing_summary[:2000],
                "feedback": feedback[:1000],
            },
            metadata={"app": APP_NAME, "mode": "feedback_regen", "user_id": USER_ID},
        ) as span:
            new_summary = await _run_core()
            span.update(output=new_summary[:4000])
            return new_summary

    return await _run_core()


def run_regen_sync(existing_summary: str, feedback: str) -> str:
    return asyncio.run(run_regen_async(existing_summary, feedback))


def build_layout():
    st.set_page_config(page_title="Agentic Claims Pipeline", page_icon="🧠", layout="wide")
    st.title("🧠 Agentic Claim Summarisation (with Q&A, Feedback & Langfuse)")

    col_left, col_mid, col_right = st.columns([1.3, 2.2, 2.5])

    with col_left:
        st.markdown("### How it works")
        st.markdown(
            """
Pipeline (ADK SequentialAgent):

1. 🧲 Data Ingestion Agent – cleans raw claim text  
2. 🍀 Parser Agent – extracts structured fields  
3. 📄 Summary Agent – drafts narrative summary  
4. 📑 Formatter Agent – formats into sections  
5. 🧮 Critic Agent – spots issues / gaps  
6. 🔁 Regeneration Agent – rewrites final summary  

If LANGFUSE_PUBLIC_KEY and LANGFUSE_SECRET_KEY are set,
all runs are traced to Langfuse for observability.
"""
        )

    with col_mid:
        st.markdown("#### Upload claim document")
        uploaded_file = st.file_uploader(
            "Drag and drop file here",
            type=["pdf", "docx", "txt"],
            help="Limit ~200MB per file · PDF, DOCX, TXT",
        )

        extracted_text = ""
        if uploaded_file is not None:
            extracted_text = parse_uploaded_file(uploaded_file)
            st.session_state[UPLOADED_DOC_KEY] = extracted_text

            st.markdown(
                f"**{uploaded_file.name}**  \n"
                f"{uploaded_file.size / 1024:.1f} KB"
            )

            with st.expander("Preview extracted text", expanded=True):
                st.caption("Extracted text")
                preview = extracted_text[:4000]
                st.text_area(
                    "preview",
                    value=preview,
                    height=260,
                    label_visibility="collapsed",
                )
        else:
            st.info("Upload a claim document to enable the pipeline.")
            st.session_state[UPLOADED_DOC_KEY] = ""

    with col_right:
        st.markdown("#### 🧬 Agentic Claim Summarisation")
        status_placeholders: Dict[str, st.delta_generator.DeltaGenerator] = {}
        for agent_name in STEP_LABELS.keys():
            label = STEP_LABELS[agent_name]
            box = st.container()
            box.info(f"⏳ {label}: waiting...")
            status_placeholders[agent_name] = box

        can_run = bool(st.session_state.get(UPLOADED_DOC_KEY))
        run_button = st.button(
            "Run ADK Claims Pipeline",
            type="primary",
            use_container_width=True,
            disabled=not can_run,
        )

    return run_button, status_placeholders


def main():
    require_api_key()

    run_button, status_placeholders = build_layout()

    st.markdown("---")
    st.markdown("### Final Claim Summary (ADK)")

    if run_button:
        claim_text = st.session_state.get(UPLOADED_DOC_KEY, "")
        if not claim_text:
            st.warning("Please upload a claim document first.")
            return

        with st.spinner("Running multi-agent claims pipeline..."):
            try:
                summary = run_pipeline_sync(claim_text, status_placeholders)
            except Exception as e:
                logging.exception("Error while running ADK pipeline:")
                st.error(f"Error while running ADK pipeline: {e}")
                return

            st.session_state[PIPELINE_RESULT_KEY] = summary

    final_text = st.session_state.get(PIPELINE_RESULT_KEY)
    if final_text:
        st.markdown(final_text)
    else:
        st.info("Run the pipeline to see the final claim summary here.")

    st.markdown("---")
    st.markdown("### Work with this claim")

    if not st.session_state.get(UPLOADED_DOC_KEY):
        st.info("Upload a claim document to enable Q&A and feedback.")
        return

    tab_qa, tab_feedback = st.tabs(
        ["❓ Ask questions about this claim", "✏️ Refine summary with feedback"]
    )

    claim_text = st.session_state.get(UPLOADED_DOC_KEY, "")

    with tab_qa:
        question = st.text_input(
            "Ask a question about the uploaded claim document:",
            placeholder="e.g., What is the date of loss? Who is the claimant?",
        )
        if st.button("Get answer", key="qa_button") and question:
            with st.spinner("Answering your question from the claim document..."):
                try:
                    answer = run_qa_sync(question, claim_text)
                except Exception as e:
                    logging.exception("Error while running Q&A:")
                    st.error(f"Error while running Q&A: {e}")
                else:
                    st.session_state[QA_RESULT_KEY] = answer

        qa_answer = st.session_state.get(QA_RESULT_KEY)
        if qa_answer:
            st.markdown("#### Answer")
            st.markdown(qa_answer)

    with tab_feedback:
        if not final_text:
            st.info("Run the pipeline first to generate a summary, then you can refine it.")
        else:
            st.markdown("Current summary (read-only):")
            st.markdown(final_text)

            feedback = st.text_area(
                "Provide feedback to change the summary:",
                placeholder="e.g., Emphasise liability, clarify coverage issues, add more detail on damages...",
                height=160,
            )

            if st.button("Regenerate summary with feedback", key="regen_button") and feedback.strip():
                with st.spinner("Regenerating summary with your feedback..."):
                    try:
                        new_summary = run_regen_sync(final_text, feedback)
                    except Exception as e:
                        logging.exception("Error while running regeneration:")
                        st.error(f"Error while running regeneration: {e}")
                    else:
                        st.session_state[PIPELINE_RESULT_KEY] = new_summary
                        st.success("Summary regenerated.")
                        st.markdown("#### Updated Summary")
                        st.markdown(new_summary)


if __name__ == "__main__":
    main()
